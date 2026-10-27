"""Writes the Excel revision matrix (one sheet, "Revision Matrix").

The workbook is only a view of what is in SQLite. It is never read back, so editing the file by hand changes
nothing; export again and it matches the database.
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import action_logger as alog
import config
import memory_manager as mm

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
VERDICT_FILL = {
    "Addressed": PatternFill("solid", fgColor="C6EFCE"),
    "Partial": PatternFill("solid", fgColor="FFEB9C"),
    "Unverified": PatternFill("solid", fgColor="F8CBAD"),
    "Ask Author": PatternFill("solid", fgColor="D9B3FF"),
}

MATRIX_COLUMNS = [
    ("Paper ID", 14), ("Review cycle ID", 16), ("Document version", 11), ("Comment ID", 10), ("Comment status", 14),
    ("Reviewer comment text (as judged)", 48), ("Author resolution claim", 28), ("Claim assessment", 14),
    ("Agent-suggested verdict", 16), ("Author-approved verdict", 18), ("Approval state", 16),
    ("Supporting evidence (verbatim quotes)", 60), ("Section and page references", 34),
    ("Missing evidence / blockers / contradictions", 46), ("Recommended next action", 44), ("Next action status", 14),
    ("Follow-up question", 48), ("Previous verdict / history", 34), ("Author corrections and comments", 48),
    ("Regression flag", 11), ("Engine", 22), ("Memory mode", 14), ("Suggested at (UTC)", 20),
]


def _ref(heading, page) -> str:
    return "%s (%s)" % (heading or "no heading", "p.%d" % page if page else "page n/a")


def _verdict_history(conn, paper_id: str, comment_id: str) -> dict[int, str]:
    rows = conn.execute(
        "SELECT cy.cycle_no, v.verdict, v.is_current FROM verdicts v JOIN review_cycles cy ON cy.paper_id = v.paper_id"
        " AND cy.cycle_id = v.cycle_id WHERE v.paper_id = ? AND v.comment_id = ? ORDER BY cy.cycle_no, v.verdict_id",
        (paper_id, comment_id)).fetchall()
    out: dict[int, str] = {}
    for r in rows:
        out[r["cycle_no"]] = r["verdict"] + ("" if r["is_current"] else " (superseded)")
    return out


def matrix_rows(conn, paper_id: str) -> list[list]:
    """One row per current agent suggestion per cycle, plus rows for comments not yet evaluated or withdrawn."""
    comments = {c["comment_id"]: c for c in mm.active_comments(conn, paper_id)}
    order = list(comments)
    cycles = conn.execute("SELECT cycle_id, cycle_no, version_id FROM review_cycles WHERE paper_id=? ORDER BY cycle_no", (paper_id,)).fetchall()
    rows: list[list] = []
    last_cycle = cycles[-1] if cycles else None
    for cyc in cycles:
        covered = set()
        vs = conn.execute("SELECT * FROM verdicts WHERE paper_id=? AND cycle_id=? AND is_current=1", (paper_id, cyc["cycle_id"])).fetchall()
        vs = sorted(vs, key=lambda r: order.index(r["comment_id"]) if r["comment_id"] in order else 999)
        for v in vs:
            covered.add(v["comment_id"])
            cid = v["comment_id"]
            rev = conn.execute("SELECT * FROM comment_revisions WHERE paper_id=? AND rev_id=?", (paper_id, v["comment_rev_id"])).fetchone()
            claim = conn.execute("SELECT claim, note FROM author_claims WHERE paper_id=? AND cycle_id=? AND comment_id=? AND is_current=1",
                                 (paper_id, cyc["cycle_id"], cid)).fetchone()
            dec = conn.execute("SELECT * FROM author_decisions WHERE paper_id=? AND cycle_id=? AND comment_id=? AND is_current=1",
                               (paper_id, cyc["cycle_id"], cid)).fetchone()
            applies = bool(dec) and (dec["decision_type"] == "correct" or dec["verdict_id"] == v["verdict_id"])
            if applies:
                approved = dec["final_verdict"]
                state = "corrected by author" if dec["decision_type"] == "correct" else "approved by author"
            else:
                approved = ""
                state = "stale (re-evaluated since approval)" if dec else "PENDING author approval"
            ev = conn.execute("SELECT * FROM verdict_evidence WHERE paper_id=? AND verdict_id=? ORDER BY evidence_id",
                              (paper_id, v["verdict_id"])).fetchall()
            evidence = "\n".join('[%s] "%s"' % (_ref(e["heading"], e["page"]), e["quote"]) for e in ev) or "(no verifiable evidence found)"
            refs = "\n".join(dict.fromkeys(_ref(e["heading"], e["page"]) for e in ev)) or "n/a"
            gaps = " | ".join(x for x in (v["missing_evidence"], v["contradiction"]) if x)
            act = conn.execute("SELECT status FROM next_actions WHERE paper_id=? AND cycle_id=? AND comment_id=? AND action_code=?"
                               " ORDER BY action_id DESC LIMIT 1", (paper_id, cyc["cycle_id"], cid, v["action_code"])).fetchone()
            hist = _verdict_history(conn, paper_id, cid)
            earlier = "; ".join("C%d: %s" % (n, t) for n, t in hist.items() if n < cyc["cycle_no"]) or "none (first evaluation)"
            prev_txt = earlier
            if v["previous_verdict"]:
                prev_txt = "latest before this cycle: %s (%s). All: %s" % (v["previous_verdict"], v["previous_source"], earlier)
            notes = []
            for k in conn.execute("SELECT kind, text, created_at FROM clarifications WHERE paper_id=? AND comment_id=? AND cycle_id=? ORDER BY clar_id",
                                  (paper_id, cid, cyc["cycle_id"])):
                notes.append("%s (%s): %s" % (k["kind"], k["created_at"][:10], k["text"]))
            if dec and dec["note"]:
                notes.append("author note: " + dec["note"])
            if rev and rev["rev_no"] > 1:
                notes.append("comment corrected to revision %d%s" % (rev["rev_no"], ": " + rev["reason"] if rev["reason"] else ""))
            claim_txt = "%s%s" % (claim["claim"], " - " + claim["note"] if claim and claim["note"] else "") if claim else "none recorded"
            rows.append([
                paper_id, cyc["cycle_id"], cyc["version_id"].rsplit("-", 1)[-1], cid, "active", rev["text"] if rev else "",
                claim_txt, v["claim_assessment"] or "", v["verdict"], approved, state, evidence, refs, gaps,
                "%s: %s" % (v["action_code"], v["recommended_action"] or ""), act["status"] if act else "",
                v["follow_up_question"] or "", prev_txt, "\n".join(notes), "YES" if v["regression"] else "",
                v["engine"], v["memory_mode"], v["created_at"]])
        if cyc is last_cycle:
            for cid in order:
                if cid not in covered:
                    c = comments[cid]
                    withdrawn = c["status"] == "withdrawn"
                    rows.append([paper_id, cyc["cycle_id"], cyc["version_id"].rsplit("-", 1)[-1], cid,
                                 "WITHDRAWN" if withdrawn else "active", c["text"], "", "", "", "",
                                 "withdrawn by author" if withdrawn else "not evaluated in this cycle", "", "", "",
                                 "" if withdrawn else "Evaluate this comment against the latest cycle", "", "", "", "", "", "", "", ""])
    if not cycles:
        for cid in order:
            rows.append([paper_id, "(baseline only)", "v1", cid, comments[cid]["status"], comments[cid]["text"]] + [""] * 17)
    return rows


def export_workbook(conn, paper_id: str, out_path=None) -> dict:
    mm.require_paper(conn, paper_id)
    paths = config.ensure_dirs()
    target = paths.exports / ("%s_revision_matrix.xlsx" % paper_id)
    if out_path:
        target = type(target)(out_path)
        target.parent.mkdir(parents=True, exist_ok=True)

    rows = matrix_rows(conn, paper_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Revision Matrix"
    ws.append([name for name, _ in MATRIX_COLUMNS])
    for r in rows:
        ws.append(r)
    for i, (_, width) in enumerate(MATRIX_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        head = ws.cell(row=1, column=i)
        head.fill, head.font, head.border = HEADER_FILL, HEADER_FONT, BORDER
        head.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
            if c.column in (9, 10) and c.value in VERDICT_FILL:
                c.fill = VERDICT_FILL[c.value]
    ws.freeze_panes = "E2"
    ws.row_dimensions[1].height = 32
    if rows:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(MATRIX_COLUMNS)), ws.max_row)

    # save to a temp name first so a failed save never leaves a half written workbook
    tmp = target.with_name(target.name + ".tmp")
    wb.save(str(tmp))
    os.replace(str(tmp), str(target))
    alog.log_memory(conn, "read", paper_id, None, "export:" + target.name,
                    "workbook generated from SQLite: %d matrix rows (database not modified)" % len(rows))
    return {"path": str(target), "matrix_rows": len(rows)}


def read_matrix_summary(path) -> list[dict]:
    """Read the matrix sheet back. Only the tests and the evaluation use this."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True)
    rows = list(wb["Revision Matrix"].iter_rows(values_only=True))
    wb.close()
    return [dict(zip(rows[0], r)) for r in rows[1:]]
